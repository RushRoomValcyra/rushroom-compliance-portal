import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

spec_path, out_path = sys.argv[1], sys.argv[2]
spec = json.load(open(spec_path, encoding="utf-8"))
geo = spec["geometry_mm"]
variants = sorted(spec["variants"], key=lambda v: v["code"])
ELIGIBLE = {"S": [0, 2, 4, 5, 7], "M": [0, 1, 2, 4, 5, 6, 7], "L": [0, 1, 2, 3, 4, 5, 6, 7]}
PANELS = [  # panel, Gateway article, faces that meet a module
    ("Left side", "Left Side Panel - White (in Prepared Left Side Panel, RR-202609-DEFPG54X)", ["right"]),
    ("Mid panel", "Mid panel part not created yet (Prepared Mid Panel, RR-202609-7C4GCN8M)", ["left", "right"]),
    ("Right side", "Right Side Panel - White (in Prepared Right Side Panel, RR-202609-WWQTVTEA)", ["left"]),
]

# The planner's interior item kinds (valcyra-site partKeys.js): a condition's "contains".
KINDS = [  # kind, name, sizes it comes in, what it is
    ("drawer_set", "Drawers", "M, L", "Drawer set, two drawers"),
    ("shelf", "Shelves", "S, M, L", "Loose shelves on pins"),
    ("hanger", "Clothes rail", "S, M, L", "Rail, one or two per module"),
    ("led_ramp", "LED strip", "S, M, L", "Under the top shelf, every module"),
    ("beautybox", "BeautyBox", "L", "Integrated unit, replaces the interior"),
    ("removable_panel_back_full", "Back panel", "S, M, L", "Full-height removable back, when chosen"),
]

def presets_with(kind):
    out = []
    for v in variants:
        c = v["content"]
        has = {"drawer_set": bool(c.get("drawers")), "shelf": bool(c.get("shelves")) and not c.get("beautybox"),
               "hanger": bool(c.get("rails")), "beautybox": bool(c.get("beautybox")),
               "led_ramp": True, "removable_panel_back_full": True}[kind]
        if has: out.append(f'{v["code"]} {v["name_sv"]}')
    return ", ".join(out)

def heights(kind):
    if kind == "drawer_set": return "Drawers " + ", ".join(f"{d['bottom']}–{d['top']}" for d in geo["drawers"])
    if kind == "shelf":
        boards = sorted({s for v in variants for s in v["content"].get("shelves", [])}, key=lambda b: -geo["boards"][b])
        return "Shelves at " + ", ".join(f"{geo['boards'][b]}" for b in boards) + " (which ones depend on the interior)"
    if kind == "hanger": return "Rails at " + ", ".join(f"{k} {h}" for k, h in geo["rails"].items())
    if kind == "led_ramp": return f"Under the top fixed shelf ({geo['boards']['Fixed shelf 1']})"
    return "—"

F = "Arial"
font = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True); head = Font(name=F, size=10, bold=True, color="FFFFFF")
title = Font(name=F, size=14, bold=True)
fill_head = PatternFill("solid", fgColor="2F4858"); fill_input = PatternFill("solid", fgColor="FFF2CC"); fill_ex = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="D0D5DA"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ---- Read me
rm = wb.active; rm.title = "Read me"
rows = [
    ("Hub operations matrix — side panels", title),
    ("Which extra holes or milling each side panel needs at the hub, per item in the module next to it. Order Operations item 0039; the E&C Gateway defines, Order Operations executes (decision 0012).", font),
    ("", font),
    ("How it works", bold),
    ("1. Each row is one face of a side panel next to a module that holds one kind of item: drawers, shelves, a clothes rail, the LED strip, a BeautyBox or a back panel. A mid panel has two faces, so each item appears twice for it.", font),
    ("2. Most rows are covered by the supplier's default drilling. Mark those 'Yes' and leave the rest of the row empty.", font),
    ("3. Where the hub must add holes or milling, mark 'No' and fill in Operation, Instruction and Drawing. Leave Size at 'Any' unless the holes differ by module size; then copy the row once per size.", font),
    ("4. Each 'No' row becomes one operation on that side panel part in the E&C Gateway (Manufacturing tab), with the condition in column M: Face, Size and 'Module has'. After Resolve BOM, Order Operations looks at the module on each side of every side panel in the order and puts the matching operations on the order's work order, with the panel's position and face.", font),
    ("5. One row covers every interior that holds the item: drawers are in Lådor + häng and Lådor + hyllor alike. When the holes depend on the interior itself (for example which shelf heights), use the Gateway's Interior choice instead and say so in Notes.", font),
    ("", font),
    ("Legend", bold),
    ("Yellow cells: fill in. Grey cells: from the planner (spec/interior-variants.json, lib/experience/utils/partKeys.js), do not edit. Column M is a formula.", font),
    ("Heights are millimetres in the planner's 3D model (feet at 13.5, side bottom edge at about 20). Subtract about 20 mm to measure from the side panel's bottom edge.", font),
    ("Faces are as seen from the front: the left face of a panel meets the module to its left.", font),
    ("Article numbers are the White variants; the same operations apply to every colour unless Notes says otherwise.", font),
    ("", font),
    ("Example (not real data)", bold),
]
for i, (text, f) in enumerate(rows, 1):
    c = rm.cell(row=i, column=1, value=text); c.font = f; c.alignment = Alignment(wrap_text=True, vertical="top")
ex_head = ["Panel", "Face", "Module has", "Supplier default enough?", "Operation", "Size", "Instruction", "Drawing", "Condition"]
ex_vals = ["Mid panel", "left", "Drawers", "No", "drill", "Any", "Drill 4 runner holes Ø5×12 for the drawer pair", "DRW-0123 rev A", '{"face":"left","contains":"drawer_set"}']
r0 = len(rows) + 1
for j, (h, v) in enumerate(zip(ex_head, ex_vals), 1):
    a = rm.cell(row=r0, column=j, value=h); a.font = bold; a.border = border
    b = rm.cell(row=r0 + 1, column=j, value=v); b.font = font; b.fill = fill_ex; b.border = border; b.alignment = wrap
rm.column_dimensions["A"].width = 110
for col in "BCDEFGHI": rm.column_dimensions[col].width = 18

# ---- Matrix
mx = wb.create_sheet("Side panel matrix")
headers = ["Panel", "Gateway article", "Face", "Module has", "Item kind", "In interiors", "Where in the module (mm)",
           "Supplier default enough?", "Operation", "Size", "Instruction", "Drawing", "Condition (for the Gateway)", "Notes"]
widths = [12, 44, 7, 13, 22, 40, 50, 13, 12, 8, 44, 16, 46, 30]
for j, (h, w) in enumerate(zip(headers, widths), 1):
    c = mx.cell(row=1, column=j, value=h); c.font = head; c.fill = fill_head; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = border
    mx.column_dimensions[c.column_letter].width = w
mx.freeze_panes = "H2"
r = 2
for panel, article, faces in PANELS:
    for face in faces:
        for kind, name, sizes, _ in KINDS:
            vals = [panel, article, face, name, kind, presets_with(kind), heights(kind), None, None, "Any", None, None, None, None]
            for j, val in enumerate(vals, 1):
                c = mx.cell(row=r, column=j, value=val); c.font = font; c.border = border; c.alignment = wrap
                c.fill = fill_input if j in (8, 9, 10, 11, 12, 14) else PatternFill("solid", fgColor="F2F2F2")
            m = mx.cell(row=r, column=13)
            m.value = (f'=IF(H{r}="No","{{""face"":"""&C{r}&""","&IF(OR(J{r}="",J{r}="Any"),"","""size"":"""&J{r}&""",")'
                       f'&"""contains"":"""&E{r}&"""}}","")')
            m.fill = PatternFill("solid", fgColor="FFFFFF"); m.font = font; m.border = border; m.alignment = wrap
            r += 1
last = r - 1
dv1 = DataValidation(type="list", formula1='"Yes,No,Unsure"', allow_blank=True); mx.add_data_validation(dv1); dv1.add(f"H2:H{last}")
dv2 = DataValidation(type="list", formula1='"drill,route,cut,insert,attach,surface_treat,inspect,other"', allow_blank=True); mx.add_data_validation(dv2); dv2.add(f"I2:I{last}")
dv3 = DataValidation(type="list", formula1='"Any,S,M,L"', allow_blank=False); mx.add_data_validation(dv3); dv3.add(f"J2:J{last}")
mx.auto_filter.ref = f"A1:N{last}"
mx["H1"].comment = Comment("Yes: the supplier's predrilled panel already has every hole this item needs. No: the hub adds holes or milling.", "Order Operations")
mx["J1"].comment = Comment("Any unless the holes differ by module size. Drawers come in M and L, the BeautyBox in L only.", "Order Operations")

# ---- Interior variants
iv = wb.create_sheet("Interior variants")
ih = ["Code", "Interior", "Sizes", "Shelves (height mm)", "Rails (height mm)", "Drawers", "Module has", "Description"]
for j, h in enumerate(ih, 1):
    c = iv.cell(row=1, column=j, value=h); c.font = head; c.fill = fill_head; c.border = border
for i, v in enumerate(variants, 2):
    c = v["content"]
    sizes = ", ".join(s for s in "SML" if v["code"] in ELIGIBLE[s])
    vals = [v["code"], v["name_sv"], sizes,
            ", ".join(f"{s} {geo['boards'][s]}" for s in c.get("shelves", [])) or "—",
            ", ".join(f"{x['rail']} {geo['rails'][x['rail']]}" for x in c.get("rails", [])) or "—",
            c.get("drawers", 0),
            ", ".join(name for kind, name, _, _ in KINDS if f'{v["code"]} {v["name_sv"]}' in presets_with(kind)),
            v.get("desc_sv", "")]
    for j, val in enumerate(vals, 1):
        cell = iv.cell(row=i, column=j, value=val); cell.font = font; cell.border = border; cell.alignment = wrap
for col, w in zip("ABCDEFGH", [6, 18, 8, 40, 32, 8, 44, 60]): iv.column_dimensions[col].width = w
n = len(variants) + 3
iv.cell(row=n, column=1, value="Fixed shelf 1 sits at 2055 in every interior; bottom board at 131. Source: valcyra-site spec/interior-variants.json (geometry_mm, measured from the GLB models 2026-10-08). Size rules: docs/wardrobe-blender-spec.md §5 (BeautyBox: L only, with a neighbour on both sides).").font = font

# ---- Summary
sm = wb.create_sheet("Summary")
sh = ["Panel", "Rows", "Answered", "Supplier default", "Hub work", "Hub rows without drawing"]
for j, h in enumerate(sh, 1):
    c = sm.cell(row=1, column=j, value=h); c.font = head; c.fill = fill_head; c.border = border
M = "'Side panel matrix'"
for i, (panel, _, _) in enumerate(PANELS, 2):
    sm.cell(row=i, column=1, value=panel)
    sm.cell(row=i, column=2, value=f'=COUNTIFS({M}!$A$2:$A${last},A{i})')
    sm.cell(row=i, column=3, value=f'=COUNTIFS({M}!$A$2:$A${last},A{i},{M}!$H$2:$H${last},"Yes")+COUNTIFS({M}!$A$2:$A${last},A{i},{M}!$H$2:$H${last},"No")')
    sm.cell(row=i, column=4, value=f'=COUNTIFS({M}!$A$2:$A${last},A{i},{M}!$H$2:$H${last},"Yes")')
    sm.cell(row=i, column=5, value=f'=COUNTIFS({M}!$A$2:$A${last},A{i},{M}!$H$2:$H${last},"No")')
    sm.cell(row=i, column=6, value=f'=COUNTIFS({M}!$A$2:$A${last},A{i},{M}!$H$2:$H${last},"No",{M}!$L$2:$L${last},"")')
t = len(PANELS) + 2
sm.cell(row=t, column=1, value="Total").font = bold
for j in range(2, 7):
    col = sm.cell(row=1, column=j).column_letter
    sm.cell(row=t, column=j, value=f"=SUM({col}2:{col}{t-1})").font = bold
for row in sm.iter_rows(min_row=2, max_row=t):
    for c in row:
        c.border = border
        if c.font != bold: c.font = font
for col, w in zip("ABCDEF", [14, 8, 10, 16, 10, 24]): sm.column_dimensions[col].width = w

wb.save(out_path)
print("rows", last - 1)
